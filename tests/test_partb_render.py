"""Part B .docx structure: banner twice with unique drawing ids, fixed table rows, the
explicit no-tickets line, group layout, filename."""
from __future__ import annotations

import datetime as dt

from docx import Document
from docx.oxml.ns import qn

from reportgen.partb import render, sites as sites_mod
from reportgen.partb.bhs_glideview import extract_zqn
from tests.test_partb_extract import WE, WS, make_zqn

SUBTITLE = "Monday 19th October to Sunday 25th October 2026"
SITES = {
    "ZQN": {"name": "Queenstown Airport", "region": "New Zealand",
            "prepared_for": "Sean Whittaker, Queenstown Airport"},
    "ST_MSE": {"name": "StarTrack Dandenong", "region": "Parcels", "status": "Active"},
    "ST_CHL": {"name": "StarTrack Chullora", "region": "Parcels", "status": "Active"},
}
ENTRY = {"kind": "ticket", "ticket_id": "181026", "call_numbers": ["15308"],
         "timestamps": ["2026-10-20T13:03:00"], "reported": "Site down",
         "findings": "Server fault", "remedy": "Restarted GSS01"}


def _texts(doc) -> list[str]:
    return [p.text for p in doc.paragraphs]


def test_zqn_report(tmp_path):
    make_zqn(tmp_path / "zqn.xlsx")
    st = extract_zqn(tmp_path / "zqn.xlsx", WS, WE)
    spec = sites_mod.resolve(["ZQN"], SITES)[0]
    out = render.render(spec, SUBTITLE, {}, {"ZQN": st}, tmp_path / spec.filename,
                        tmp_path / "charts")
    assert out.name == "ZQN_Weekly_Phone_Logs_and_BHS_Statistics.docx"

    doc = Document(str(out))
    texts = _texts(doc)
    assert "ZQN Weekly Phone Logs Report & BHS Statistics" in texts
    assert ("No support tickets were logged for Queenstown Airport (ZQN) during this period."
            in texts)
    assert "Queenstown Airport — BHS Performance Review" in texts
    assert "Prepared for: Sean Whittaker, Queenstown Airport" in texts
    assert "Prepared by: Tiaan Botha, Alstef Group" in texts
    assert texts[-3:] == ["Warm regards,", "ALSTEF Group", "Support Services Team"]

    # banner on page 1 and again on the review page, every drawing id unique
    body = doc.element.body
    doc_prs = [d.get("id") for d in body.iter(qn("wp:docPr"))]
    assert len(doc_prs) == len(set(doc_prs))
    graphics = [d for d in body.iter(qn("wp:docPr")) if d.get("name") == "Graphic 17"]
    assert len(graphics) == 2
    assert len(body.findall(".//" + qn("w:pageBreakBefore"))) == 1

    kpi = [t for t in doc.tables if t.rows[0].cells[0].text == "Metric"][0]
    assert [r.cells[0].text for r in kpi.rows[1:]] == [
        "Total throughput (main line)", "Oversize line volume", "L1 flagged for re-screen",
        "L2 confirmed reject (final)", "Lost-bag events"]

    # three charts (trend, faults, lost bags) for a full ZQN week
    assert len(body.findall(".//" + qn("pic:pic"))) - 2 == 3  # minus the two logos

    exec_idx = texts.index("Executive Summary")
    assert texts.index("1. Total Throughput") - exec_idx - 1 == 5  # exactly five bullets


def test_group_report_one_section_per_site_no_bhs(tmp_path):
    spec = sites_mod.resolve(["StarTrack"], SITES)[0]
    out = render.render(spec, SUBTITLE, {"ST_MSE": [ENTRY]}, {}, tmp_path / spec.filename,
                        tmp_path / "charts")
    texts = _texts(Document(str(out)))
    assert "StarTrack Chullora ST_CHL — Support Tickets" in texts
    assert "StarTrack Dandenong ST_MSE — Support Tickets" in texts
    assert "15308 (20/10/2026 13:03)" in texts  # Part A's hand-made meta format
    assert not any("BHS Performance Review" in t for t in texts)
    assert texts.index("StarTrack Chullora ST_CHL — Support Tickets") < texts.index(
        "StarTrack Dandenong ST_MSE — Support Tickets")
