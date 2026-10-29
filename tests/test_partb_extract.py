"""Part B extractors, recalc staleness check, registry and formatting - on small synthetic
workbooks shaped like the real GlideView / BAGanalytics exports."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from reportgen.discover import DiscoveryError
from reportgen.partb import metrics, sites as sites_mod
from reportgen.partb.bhs_bag import extract_adl
from reportgen.partb.bhs_glideview import extract_chc, extract_zqn
from reportgen.partb.grid import num
from reportgen.partb.recalc import summary_looks_stale

WS, WE = dt.date(2026, 10, 19), dt.date(2026, 10, 25)


# --------------------------------------------------------------------------------------
# builders
# --------------------------------------------------------------------------------------


def _fault_sheet(wb, devices: dict[str, dict[str, int]]):
    ws = wb.create_sheet("Fault Table")
    ws.append(["Site - GlideView Reports"])
    ws.append(["Fault Table Week Ending Sunday, 25 October 2026"])
    ws.append([])
    names = sorted({n for d in devices.values() for n in d})
    ws.append([None, *names, "TOTAL"])
    for dev, counts in devices.items():
        ws.append([dev, *[str(counts.get(n, 0)) for n in names], "0"])
    ws.append(["TOTAL", *["999"] * len(names), "0"])


def make_zqn(path: Path, days: int = 7, t101_per_row: int = 10, summary_t101: int | None = None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rates = wb.create_sheet("System Throughput Rates")
    rates.append(["Queenstown International Airport - GlideView Reports"])
    rates.append(["System Throughput Rates Week Ending Sunday, 25 October 2026"])
    rates.append([])
    rates.append(["UTCTime", "LocalTime", "OS104", "T101", "T114"])
    first = WE - dt.timedelta(days=days - 1)
    total = 0
    for d in range(days):
        for hour in (1, 12, 24):  # 24 -> the next day's 00:00 row, which belongs to today
            local = dt.datetime.combine(first + dt.timedelta(days=d), dt.time()) + dt.timedelta(hours=hour)
            rates.append([local - dt.timedelta(hours=13), local, 0, t101_per_row, 0])
            total += t101_per_row

    s = wb.create_sheet("SUMMARY")
    for row in (
        ["Queenstown International Airport - GlideView Reports"],
        ["Summary Report Week Ending Sunday, 25 October 2026"],
        ["Main Outbound System"], ["T101", summary_t101 if summary_t101 is not None else total],
        ["T114", "1,234"], ["Oversize Outbound System"], ["OS104", 50],
        ["L1 Clear Count", 60], ["L1 Reject Bag Count", 30], ["L1 No Scan Count", 9],
        ["L1 Timeout Count", 1], ["L2 Clear Count", 25], ["L2 Reject Bag Count", 3],
        ["L2 Timeout Count", 2],
        ["Tracking Statistics"], ["Zone", "Tracked Bag Count", "Lost Bag Count"],
        ["T109PC", 0, 0], ["T110PC", 1000, 4], ["T111PC", 1000, 0], ["T112PC", 1000, 1],
        ["T1DV1PD", 1000, 2], ["T1DV1PU", 1000, 1], ["T110APC", 1000, 2], ["TOTAL", 10],
    ):
        s.append(row)
    _fault_sheet(wb, {"T101": {"Dieback Status": 500, "Bag Jam Fault": 4},
                      "T102": {"Bag Jam Fault": 3, "Emergency Stop Fault": 2}})
    wb.save(path)
    return total


CHC_ZONES = [("TC1 & TC2 prior to ME", 10), ("Domestic CBS", 40), ("International CBS", 30),
             ("Exiting ME", 5), ("Between ME and L4T01", 1), ("S1 Sort Line", 2),
             ("S2 Sort Line", 8), ("Level 4 Line", 4)]


def make_chc(path: Path, tracking_total: int = 100):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    s = wb.create_sheet("SUMMARY")
    s.append(["Christchurch International Airport - GlideView Reports"])
    s.append(["SORT SYSTEM", 2000])
    s.append(["OOG Total", 150])
    lines = [(1000, 600, 5, 10, 3), (900, 500, 4, 9, 2), (800, 450, 3, 8, 1), (700, 400, 2, 7, 0)]
    for i, (scr, clr, ns, rej, to) in enumerate(lines, start=1):
        s.append([f"X1X0{i}00 EDS"])
        for label, v in (("Bags Screened", scr), ("Level 1 Clear", clr),
                         ("Level 1 Clear %", clr / scr), ("No Scans", ns),
                         ("Level 2 Reject", rej), ("Level 2 Timeout", to)):
            s.append([f"   {label}", v])
    s.append(["SCREENING TOTAL", 3400])
    s.append(["TRACKING STATS"])
    for z, n in CHC_ZONES:
        s.append([z, n, 0.01, "Note"])
    s.append(["TOTAL", tracking_total])
    s.append(["Percentage*", tracking_total / 2000, "* = percentage of total bags entering sort system"])

    sec = wb.create_sheet("Security Screening Statistics")
    sec.append(["Level 1 EDS Decisions"])
    sec.append([None, "Level 1 Clear Count", "Level 1 Reject Count", "Level 1 Timeout Count",
                "Level 1 No Scan Count"])
    for i, (scr, clr, ns, _, _) in enumerate(lines, start=1):
        sec.append([f"X1X0{i}00XR", clr, scr - clr - ns, 0, ns])
    sec.append(["Level 2 EDS Decisions"])
    sec.append([None, "Level 2 Clear Count", "Level 2 Reject Count", "Level 2 Timeout Count"])
    for i, (_, _, _, rej, to) in enumerate(lines, start=1):
        sec.append([f"X1X0{i}00XR", 100, rej, to])
    _fault_sheet(wb, {"M1": {"Emergency Stop Fault": 20, "ME Waiting Bag Warning": 30},
                      "M2": {"Motor Isolator Fault": 5, "Manual Mode Active": 1}})
    wb.save(path)


def make_adl(path: Path, eds_zero: bool = False):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    s = wb.create_sheet("Summary")
    for row in (
        ["BAGview Reports | Summary"], ["FILTERS (2)"],
        ["Date: LAST 7 DAYS: 10/19/2026 - 10/25/2026"],
        [None, None, None, "Total Check-in Bags (In Gauge)", None, None, None, None,
         "Total Check-in Bags (Oversize)"],
        [None, None, None, "59,302", None, None, None, None, "3,339"],
        [None, None, None, "Total Recircs plus Check-in Bags (In Gauge)"],
        [None, None, None, "65,802"],
        [None, None, None, "Total Make-up Output"], [None, None, None, "59,189"],
        [None, None, None, "Total Bags Processed"], [None, None, None, "62,641"],
        [None, None, None, "Lost Track Percentage"], [None, None, None, "6.28%"],
        [None, None, None, "Tracking Rate"], [None, None, None, "93.72%"],
        [None, None, None, "Total Lost Bags"], [None, None, None, "4,134"],
        [None, None, None, "Zone", None, None, None, None, "Lost Count"],
        [None, None, None, "XT1", None, None, None, None, "3,000"],
        [None, None, None, "SL1", None, None, None, None, "1,134"],
        [None, None, None, "Collector conveyor downtime"],
    ):
        s.append(row)
    f = wb.create_sheet("Faults List")
    for _ in range(8):
        f.append([])
    f.append([None, "Alarm Area ", "Alarms Equipment", "Alarm", "Alarms Count"])
    f.append([None, "OB2", " ABD121 ", " EStop Fault", 17])
    f.append([None, None, None, " Heartbeat Fault", 1132])
    f.append([None, None, None, "Total", 1149])
    f.append([None, None, " ABD123 ", " EStop Fault", 3])
    f.append([None, None, None, " Bag Waiting Alarm", "1,339"])
    e = wb.create_sheet("EDS Statistics ")
    e.append(["Date: LAST 7 DAYS: 10/19/2026 - 10/25/2026"])
    e.append(["Level", "Component", "Error", "Fail", "Pass", "Pass (%)", "Total"])
    v = "0.00" if eds_zero else "12.00"
    e.append(["Level 2", "OS1X1", "0.00", v, "0.00", "0.97", v])
    e.append([None, "XT1X1", "0.00", "0.00", "0.00", "0.97", "0.00"])
    wb.save(path)


# --------------------------------------------------------------------------------------
# tests
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("raw,expected", [
    (" 59,302 ", 59302.0), ("6.28%", 6.28), ("0.00", 0.0), (17, 17.0), ("", None),
    (None, None), ("n/a", None),
])
def test_num_parses_export_text(raw, expected):
    assert num(raw) == expected


def test_zqn_full_week(tmp_path):
    total = make_zqn(tmp_path / "zqn.xlsx")
    st = extract_zqn(tmp_path / "zqn.xlsx", WS, WE)
    assert (st.main_throughput, st.oversize, st.lost_total) == (total, 50, 10)
    assert (st.l1_flagged, st.l1_denominator, st.l2_reject, st.l2_timeouts) == (30, 100, 3, 2)
    assert st.lost_pct == pytest.approx(100 * 10 / 6000)
    assert st.full_week and len(st.trend) == 7
    assert all(n == 30 for _, n in st.trend)  # 00:00 row lands on the previous day
    assert st.excluded == [("Dieback Status", 500, "a heartbeat/status counter, not a fault")]
    assert st.faults == [("Bag Jam Fault", 7), ("Emergency Stop Fault", 2)]  # TOTAL row/col ignored
    assert "1,234" in st.main_note


def test_zqn_single_day_has_no_trend(tmp_path):
    make_zqn(tmp_path / "zqn.xlsx", days=1)
    st = extract_zqn(tmp_path / "zqn.xlsx", WS, WE)
    assert st.single_day and st.trend is None and "single day" in st.trend_note
    assert any(f.startswith("!! export covers only 1 day") for f in st.flags)
    assert metrics.scope(st) == "on Sunday 25 October 2026 only"


def test_zqn_series_disagreeing_with_summary_hard_stops(tmp_path):
    make_zqn(tmp_path / "zqn.xlsx", summary_t101=1)
    with pytest.raises(DiscoveryError, match="cross-check failed"):
        extract_zqn(tmp_path / "zqn.xlsx", WS, WE)


def test_chc(tmp_path):
    make_chc(tmp_path / "chc.xlsx")
    st = extract_chc(tmp_path / "chc.xlsx", WS, WE)
    assert st.main_throughput == 3400
    assert st.l1_flagged == 3400 - (600 + 500 + 450 + 400) - (5 + 4 + 3 + 2)
    assert (st.l2_reject, st.l2_timeouts) == (34, 6)
    assert (st.lost_total, st.lost_pct) == (100, pytest.approx(5.0))
    assert st.lost_breakdown == CHC_ZONES
    assert st.pm_testing == (26, 56)
    assert st.trend is None and "no 5-minute" in st.trend_note
    assert st.secondary[0][1] == "2,000 bags (C1+C2+C3+Tranship)"


def test_chc_zone_total_mismatch_hard_stops(tmp_path):
    make_chc(tmp_path / "chc.xlsx", tracking_total=99)
    with pytest.raises(DiscoveryError, match="tracking zones"):
        extract_chc(tmp_path / "chc.xlsx", WS, WE)


def test_adl_text_numbers_exclusions_and_withheld(tmp_path):
    make_adl(tmp_path / "adl.xlsx")
    st = extract_adl(tmp_path / "adl.xlsx", WS, WE)
    assert (st.main_throughput, st.oversize, st.lost_total, st.lost_pct) == (59302, 3339, 4134, 6.28)
    assert st.lost_breakdown == [("XT1", 3000), ("SL1", 1134)]
    assert st.faults == [("Bag Waiting Alarm", 1339), ("EStop Fault", 20)]
    assert st.excluded[0][:2] == ("Heartbeat Fault", 1132)
    assert st.reject_withheld and "reconciled" in st.reject_withheld
    assert st.l1_flagged is None and st.l2_reject is None
    rows = metrics.table_rows(st)
    assert ("L1 / L2 reject figures", "Provisional - withheld", "See note below the table") in rows


def test_adl_all_zero_eds_is_a_data_problem_not_zero_rejects(tmp_path):
    make_adl(tmp_path / "adl.xlsx", eds_zero=True)
    st = extract_adl(tmp_path / "adl.xlsx", WS, WE)
    assert "not 0% rejects" in st.reject_withheld


def test_adl_wrong_week_hard_stops(tmp_path):
    make_adl(tmp_path / "adl.xlsx")
    with pytest.raises(DiscoveryError, match="not the reporting week"):
        extract_adl(tmp_path / "adl.xlsx", dt.date(2026, 10, 12), dt.date(2026, 10, 18))


def test_summary_looks_stale(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SUMMARY"
    ws["A1"], ws["B1"], ws["B2"] = "Bags", 5, "=B1*2"  # saved by openpyxl: no cached value
    wb.save(tmp_path / "stale.xlsx")
    assert summary_looks_stale(tmp_path / "stale.xlsx")

    ws["B2"] = 10
    wb.save(tmp_path / "plain.xlsx")
    assert not summary_looks_stale(tmp_path / "plain.xlsx")


def test_registry_groups_and_unknown_site():
    sites = {
        "YVR": {"name": "Vancouver", "region": "Canada", "status": "Active"},
        "YQM": {"name": "Moncton", "region": "Canada", "status": "Inactive"},
        "YXS": {"name": "Prince George", "region": "Canada", "status": "Inactive"},
        "ST_MSE": {"name": "StarTrack Dandenong", "region": "Parcels", "status": "Active"},
        "ZQN": {"name": "Queenstown Airport", "prepared_for": "Sean"},
    }
    specs = sites_mod.resolve(["canadian sites", "StarTrack", "zqn"], sites, {"YQM"})
    assert [s.key for s in specs] == ["Canadian Sites", "StarTrack", "ZQN"]
    assert [s.code for s in specs[0].sites] == ["YQM", "YVR"]  # inactive shown only with entries
    assert specs[0].filename == "Canadian_Sites_Weekly_Phone_Logs_and_BHS_Statistics.docx"
    assert specs[2].sites[0].prepared_for == "Sean"
    with pytest.raises(DiscoveryError, match="one-off"):
        sites_mod.resolve(["AKL"], sites)


def test_week_range_labels():
    assert metrics.week_range(WS, WE) == "Monday 19 - Sunday 25 October 2026"
    assert metrics.week_range(dt.date(2026, 9, 28), dt.date(2026, 10, 4)) == \
        "Monday 28 September - Sunday 4 October 2026"
