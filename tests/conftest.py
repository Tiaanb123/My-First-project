from __future__ import annotations

import datetime as dt
import shutil
from pathlib import Path

import openpyxl
import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent


def make_phonelogs(path: Path, calls, stats_lines: list[str] | None = None) -> None:
    """Build a phone-log workbook.

    `calls` items are either a `datetime` (location defaults to ZQN_APAC) or a dict with
    keys number / entered / location_id / comments / transferred_to / inquiry. `entered`
    may also be a `"DD/MM/YYYY HH:MM:SS"` string - the real export mixes both.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    stats = wb.create_sheet("Phone Stats")
    for i, line in enumerate(stats_lines or ["placeholder"], start=1):
        stats.cell(row=i, column=1, value=line)
    logs = wb.create_sheet("Phone Logs")
    logs.append(["Call Number", "Call Type", "Comments", "Operator", "Time Entered",
                 "Location of the caller", "Who did you transfer the call to", "Inquiry"])
    for i, item in enumerate(calls, start=15000):
        if isinstance(item, dt.datetime):
            item = {"entered": item}
        when = item.get("entered")
        logs.append([
            item.get("number", i), "Divert",
            item.get("comments",
                     f"E-mailed on {when:%d-%b-%Y} by Op" if isinstance(when, dt.datetime) else ""),
            "Op", when, item.get("location_id", "ZQN_APAC"),
            item.get("transferred_to", "Divert #2"), item.get("inquiry", ""),
        ])
    wb.create_sheet("Data")
    wb.save(path)


def make_glideview(path: Path) -> None:
    openpyxl.Workbook().save(path)


def make_baganalytics(path: Path, start: dt.date, end: dt.date) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Summary")
    ws["A1"] = "BAGview Reports | Summary"
    ws["A2"] = "FILTERS (2)"
    ws["A3"] = (
        f"Date: LAST 7 DAYS: {start.month}/{start.day}/{start.year} - "
        f"{end.month}/{end.day}/{end.year}"
    )
    wb.save(path)


def make_carlsource(path: Path, tickets: list[dict]) -> None:
    """Build a CarlSource 'WO Report' workbook (the export's real, tabular shape).

    `tickets`: dicts with id, title, client_code, client_name, status, begin
    (`"DD/MM/YYYY HH:MM:SS"`), call_reference, and either `comments` directly or
    reported/findings/remedy/freeform to assemble it.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("WO Report")
    ws.append(["Filter criteria"])
    ws.append(["[Lines selected ]"])
    ws.append(["Name", "WO", "Call Reference", "Status code", "Start date",
               "Client Code", "Client name", "Comments"])
    for t in tickets:
        comments = t.get("comments")
        if comments is None:
            comments = (f"Call Reference: {t.get('call_reference', '')}\nHelp:\n\n"
                        f"Reported: {t.get('reported', '')}\n"
                        f"Findings: {t.get('findings', '')}\n"
                        f"Remedy: {t.get('remedy', '')}\n")
            if t.get("freeform"):
                comments += f"\n24/08:\n{t['freeform']}\n"
        ws.append([
            t["title"], t["id"], t.get("call_reference", ""), t.get("status", ""),
            t.get("begin", ""), t.get("client_code", ""), t.get("client_name", ""), comments,
        ])
    ws.append(["Lines number", float(len(tickets))])
    wb.save(path)


@pytest.fixture
def dropbox(tmp_path: Path) -> Path:
    """A well-formed DropBox for the week Mon 24 Aug - Sun 30 Aug 2026."""
    root = tmp_path / "DropBox"
    (root / "PartA").mkdir(parents=True)
    (root / "PartB").mkdir(parents=True)

    make_phonelogs(
        root / "PartA" / "Phonelogs - 24.08 to 30.08.2026.xlsx",
        [dt.datetime(2026, 8, 24, 1, 35), dt.datetime(2026, 8, 27, 9, 20),
         dt.datetime(2026, 8, 30, 22, 5)],
    )
    make_carlsource(root / "PartA" / "Carlsource - 24.08 to 30.08.2026.xlsx", [])

    shutil.copytree(REPO_ROOT / "knowledge", tmp_path / "knowledge")
    return root
