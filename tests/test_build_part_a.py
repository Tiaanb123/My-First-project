from __future__ import annotations

import datetime as dt
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

from reportgen import build_part_a, chart_pie, parse_carlsource, parse_phonelogs
from reportgen.bundle import build_part_a as make_bundle
from reportgen.link import link
from reportgen.regions import Regions
from tests.conftest import make_carlsource, make_phonelogs

REPO_ROOT = Path(__file__).resolve().parent.parent
WEEK = {"start": "2026-09-07", "end": "2026-09-13",
        "subtitle": "Monday 7th September to Sunday 13th September 2026"}


@pytest.fixture(scope="module")
def report(tmp_path_factory) -> Path:
    tmp = tmp_path_factory.mktemp("build")
    make_phonelogs(tmp / "pl.xlsx", [
        {"number": 15223, "entered": "11/09/2026 08:43:46", "location_id": "ADL_APAC",
         "comments": "E-mailed on 11-Sep-2026 at 08:45 by Op_x000D_\nCaller could not log in"},
        {"number": 15225, "entered": dt.datetime(2026, 9, 11, 16, 34, 20), "location_id": "ZQN_APAC"},
        {"number": 15230, "entered": "12/09/2026 11:55:10", "location_id": "ZQN_APAC"},
        {"number": 15231, "entered": "12/09/2026 12:37:00", "location_id": "ZQN_APAC"},
        {"number": 15234, "entered": "13/09/2026 18:19:00", "location_id": "AMX_APAC"},
    ])
    make_carlsource(tmp / "cs.xlsx", [
        {"id": "178369", "title": "ADL", "client_code": "ADL_APAC", "begin": "11/09/2026 08:43:00",
         "call_reference": "15223", "reported": "Could not open the web UI.", "remedy": "Switched window."},
        {"id": "178565", "title": "ZQN", "client_code": "ZQN_APAC", "begin": "11/09/2026 16:34:00",
         "call_reference": "15225", "reported": "Motor feedback fault."},
        {"id": "178566", "title": "ZQN2", "client_code": "ZQN_APAC", "begin": "12/09/2026 11:55:00",
         "call_reference": "15230. 15231", "comments": "Reported:\nRemedy: xx/09:\n"},
        {"id": "178562", "title": "AMX", "client_code": "AMX_APAC", "begin": "13/09/2026 18:29:00",
         "call_reference": "15234", "reported": "Motor failure."},
        {"id": "178180", "title": "ZNE", "client_code": "ZNE_APAC", "begin": "07/09/2026 20:00:00",
         "call_reference": "direct email", "reported": "Bags stopped at TC107."},
    ])
    regions = Regions.from_yaml(REPO_ROOT / "knowledge")
    phone = parse_phonelogs.parse(tmp / "pl.xlsx")
    res = link(phone.calls, parse_carlsource.parse(tmp / "cs.xlsx"), regions)
    bundle = make_bundle(WEEK, phone, res, regions)
    bundle_path = tmp / "bundle.json"
    bundle_path.write_text(json.dumps(bundle), encoding="utf-8")
    out, _ = build_part_a.build(bundle_path, tmp / "out")
    return out


def _texts(report: Path) -> list[str]:
    return [p.text.strip() for p in Document(str(report)).paragraphs]


def test_calls_by_site_sorted_desc(report):
    bundle = json.loads((report.parent.parent / "bundle.json").read_text(encoding="utf-8"))
    data = chart_pie.calls_by_site(bundle)
    assert data[0] == ("ZQN_APAC", 3)
    assert sum(n for _, n in data) == bundle["part_a"]["totals"]["calls"] == 5
    assert data == sorted(data, key=lambda kv: (-kv[1], kv[0]))


def test_report_filename(report):
    assert report.name == "Weekly_Support_Tickets_Report_07.09.2026_to_13.09.2026.docx"


def test_report_text_and_counts(report):
    texts = _texts(report)
    joined = "\n".join(texts)
    assert "Total Amount of Calls: 5" in texts
    assert "Total Amount of Direct Contacts: 1" in texts
    headings = [t for t in texts if t.endswith((" Call", " Calls"))]
    assert headings == [
        "Africa: 0 Calls", "Australia: 2 Calls", "Bermuda: 0 Calls",
        "Canada: 0 Calls", "Jamaica: 0 Calls", "New Zealand: 3 Calls",
        "United States: 0 Calls", "Parcels: 1 Call",
        "New Zealand: 1 Call",  # Parcels' nested country sub-heading (Aramex is NZ, not AU)
    ]
    assert texts.index("Parcels: 1 Call") + 1 == texts.index("New Zealand: 1 Call")
    assert texts[-3:] == ["Warm regards,", "ALSTEF Group", "Support Services Team"]
    # string-timestamp calls matched their tickets: no Direct Contact, no empty "()" stamp
    assert "15223 (11/09/2026 08:43)" in texts
    assert "15230 (12/09/2026 11:55) & 15231 (12/09/2026 12:37)" in texts
    assert "Direct Contact (" not in joined and "()" not in joined
    assert "_x000D_" not in joined
    assert "No detail recorded in WO 178566." in joined


def test_region_summary_is_a_three_column_table(report):
    grid = Document(str(report)).tables[-1]
    assert [[c.text for c in r.cells] for r in grid.rows] == [
        ["Africa: 0", "Australia: 2", "Bermuda: 0"],
        ["Canada: 0", "Jamaica: 0", "New Zealand: 3"],
        ["United States: 0", "Parcels: 1", ""],
    ]


def test_region_summary_table_is_full_width_and_even(report):
    with zipfile.ZipFile(report) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert '<w:tblW w:type="pct" w:w="5000"/>' in xml
    assert xml.count('<w:tcW w:type="pct" w:w="1666"/>') == 9  # 3x3 grid, evenly split


def test_paragraphs_have_no_space_after(report):
    """'Remove Space After Paragraph' - every generated paragraph is tight; visible gaps
    come only from the deliberate loose spacers, not per-paragraph style defaults."""
    doc = Document(str(report))
    for p in doc.paragraphs:
        if p.text.strip():
            assert p.paragraph_format.space_after == Pt(0), p.text


def test_loose_spacer_between_empty_region_headings(report):
    texts = _texts(report)
    # Bermuda has no tickets: heading -> loose spacer -> next heading, no tight override
    doc = Document(str(report))
    bermuda = texts.index("Bermuda: 0 Calls")
    spacer = doc.paragraphs[bermuda + 1]
    assert spacer.text == ""
    assert spacer.paragraph_format.space_after is None


def test_ticket_fields_are_real_bullets(report):
    with zipfile.ZipFile(report) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert xml.count("<w:numPr>") >= 6  # every Reported/Remedy/... paragraph


def test_report_has_native_chart(report):
    with zipfile.ZipFile(report) as z:
        names = z.namelist()
        assert "word/charts/chart1.xml" in names
        assert "word/embeddings/pie.xlsx" in names
        ct = z.read("[Content_Types].xml").decode("utf-8")
        assert "drawingml.chart+xml" in ct and 'Extension="xlsx"' in ct
        doc = z.read("word/document.xml").decode("utf-8")
        assert "<c:chart" in doc and "__PIE_CHART__" not in doc
        chart = z.read("word/charts/chart1.xml").decode("utf-8")
        assert "<c:v>ZQN_APAC</c:v>" in chart
        assert "<c:dPt>" not in chart
        assert "Weekly Call Logs 13/09/2026" in chart


def test_chart_fills_the_page_content_width(report):
    doc = Document(str(report))
    section = doc.sections[0]
    content_width = section.page_width - section.left_margin - section.right_margin
    with zipfile.ZipFile(report) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    import re
    extents = re.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"/>', xml)
    assert extents[-1][0] == str(int(content_width))
