"""Render one Part B document (spec sections 7-8, 10-11).

Built on `docx_base.open_report()` (the cloned letterhead - never rebuilt in code); the
ticket blocks reuse Part A's `add_site_heading` / `add_entry` unchanged so both reports
print tickets identically. Part B's own headings/table/notes follow spec section 10 and
are local to this module so the Part A output doesn't move.
"""
from __future__ import annotations

import copy
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .. import docx_base
from . import charts, metrics, text
from .sites import ReportSpec, SiteInfo
from .stats import BhsStats

HEADING_FONT = "Franklin Gothic Medium"
BODY_FONT = "Aptos"
NAVY = RGBColor(0x0B, 0x25, 0x45)
SUB_NAVY = RGBColor(0x11, 0x2F, 0x51)
SUBTITLE_GREY = RGBColor(0x59, 0x59, 0x59)
META_GREY = RGBColor(0x76, 0x76, 0x76)
PREPARED_BY = "Tiaan Botha, Alstef Group"
CHART_WIDTH = Inches(6.3)


# --------------------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------------------


def _font(run, name: str):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), name)
    return run


def _run(p, txt: str, size: float = 11, bold: bool = False, italic: bool = False,
         colour: RGBColor | None = None, font: str = BODY_FONT):
    r = _font(p.add_run(txt), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if colour is not None:
        r.font.color.rgb = colour
    return r


def _flat(p, before: float = 0, after: float = 0):
    """Zero the letterhead Normal style's 0.5" indents and 18pt space-after."""
    pf = p.paragraph_format
    pf.left_indent = pf.right_indent = Pt(0)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    return p


def _para(doc, txt: str = "", after: float = 0, **kw):
    p = _flat(doc.add_paragraph(), after=after)
    if txt:
        _run(p, txt, **kw)
    return p


def _heading(doc, txt: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    _flat(p, before=13, after=6)
    r = _run(p, txt, size=15 if level == 1 else 14, bold=True,
             colour=NAVY if level == 1 else SUB_NAVY, font=HEADING_FONT)
    r.font.all_caps = False
    return p


def _note(doc, txt: str):
    """Italic grey 10pt - data-scope / methodology notes."""
    return _para(doc, txt, size=10, italic=True, colour=META_GREY)


def _bullet(doc, txt: str):
    p = docx_base._bullet(doc, None, txt)
    for r in p.runs:
        _font(r, BODY_FONT).font.size = Pt(11)
    return p


def _picture(doc, path: Path | None):
    if path is None:
        return None
    p = _flat(doc.add_paragraph(), before=4, after=4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=CHART_WIDTH)
    return p


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _borders(table, colour: str = "BFBFBF") -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        for attr, val in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", colour)):
            edge.set(qn(f"w:{attr}"), val)
        borders.append(edge)
    tbl_pr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 70), ("bottom", 70), ("left", 100), ("right", 100)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(w))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tbl_pr.append(mar)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")


def _kpi_table(doc, rows: list[tuple[str, str, str]]):
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.autofit = False
    _borders(table)
    widths = (Inches(2.2), Inches(2.2), Inches(2.9))
    for r, values in enumerate([("Metric", "Value", "Note"), *rows]):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.width = widths[c]
            p = _flat(cell.paragraphs[0])
            if r == 0:
                _shade(cell, "0B2545")
                _run(p, value, size=10, bold=True, colour=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if r % 2 == 0:
                    _shade(cell, "F2F2F2")
                _run(p, value, size=10)
    return table


# --------------------------------------------------------------------------------------
# letterhead banner (cloned, never rebuilt)
# --------------------------------------------------------------------------------------


def _ensure_default_size(doc) -> None:
    """Spec 11 'docDefaults trap': the banner anchors need a 12pt default run size."""
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr_default = defaults.find(qn("w:rPrDefault")) if defaults is not None else None
    rpr = rpr_default.find(qn("w:rPr")) if rpr_default is not None else None
    if rpr is None:
        return
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), "24")


def _banner_snapshot(doc) -> list:
    """Deep copies of the banner block (everything in the body before our content)."""
    return [copy.deepcopy(el) for el in doc.element.body if el.tag != qn("w:sectPr")]


def _repeat_banner(doc, snapshot: list) -> None:
    """Append the banner again, starting on a new page. Each copied drawing gets a fresh
    unique wp:docPr id; Word's per-paragraph ids are dropped so they can't collide."""
    used = {int(d.get("id")) for d in doc.element.body.iter(qn("wp:docPr"))
            if (d.get("id") or "").isdigit()}
    next_id = max(used, default=0) + 1
    w14 = "http://schemas.microsoft.com/office/word/2010/wordml"
    wp14 = "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing"
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    for i, el in enumerate(snapshot):
        el = copy.deepcopy(el)
        for d in el.iter(qn("wp:docPr")):
            d.set("id", str(next_id))
            next_id += 1
        for node in el.iter():
            for attr in (f"{{{w14}}}paraId", f"{{{w14}}}textId", f"{{{wp14}}}anchorId",
                         f"{{{wp14}}}editId"):
                node.attrib.pop(attr, None)
        if i == 0:
            if el.tag == qn("w:p"):
                ppr = el.find(qn("w:pPr"))
                if ppr is None:
                    ppr = OxmlElement("w:pPr")
                    el.insert(0, ppr)
                ppr.insert(0, OxmlElement("w:pageBreakBefore"))
            else:  # banner starts with a table - put a page break in front of it
                brk = OxmlElement("w:p")
                r = OxmlElement("w:r")
                b = OxmlElement("w:br")
                b.set(qn("w:type"), "page")
                r.append(b)
                brk.append(r)
                sect.addprevious(brk)
        sect.addprevious(el)


# --------------------------------------------------------------------------------------
# sections
# --------------------------------------------------------------------------------------


def _tickets(doc, site: SiteInfo, entries: list[dict]) -> None:
    _heading(doc, f"{site.name} {site.code} — Support Tickets", 1)
    if not entries:
        _para(doc, f"No support tickets were logged for {site.name} ({site.code}) during "
                   f"this period.", italic=True, colour=SUBTITLE_GREY)
        return
    docx_base.add_site_heading(doc, site.name, site.code)
    for entry in entries:
        docx_base.add_entry(doc, entry)
        docx_base._spacer(doc)


def _bhs_table(doc, site: SiteInfo, st: BhsStats) -> None:
    _heading(doc, f"{site.name} — Weekly BHS Statistics", 1)
    _note(doc, metrics.source_line(st))
    _para(doc)
    _kpi_table(doc, metrics.table_rows(st))
    if st.limitations:
        _para(doc)
        _note(doc, "Standing limitations of this export: " + " ".join(st.limitations))
    if st.recalculated:
        _note(doc, text.RECALC_NOTE)


def _review(doc, site: SiteInfo, st: BhsStats, entries: list[dict], banner: list,
            chart_dir: Path) -> None:
    _repeat_banner(doc, banner)
    _heading(doc, f"{site.name} — BHS Performance Review", 1)
    _para(doc, "Period covered: " + metrics.review_period(st), italic=True)
    if site.prepared_for:
        line = "Prepared for: " + site.prepared_for
        if site.prepared_for_cc:
            line += f" (CC: {site.prepared_for_cc})"
        _para(doc, line)
    _para(doc, "Prepared by: " + PREPARED_BY)

    _heading(doc, "Executive Summary", 2)
    for b in text.exec_summary(st, entries):
        _bullet(doc, b)

    _heading(doc, "1. Total Throughput", 2)
    _para(doc, text.throughput(st))
    _picture(doc, charts.trend_chart(st, chart_dir))

    _heading(doc, "2. Reject Bags", 2)
    for para in text.rejects(st):
        _para(doc, para, after=6)

    _heading(doc, "3. Fault Subsections", 2)
    _para(doc, text.faults(st))
    _picture(doc, charts.fault_chart(st, chart_dir))

    _heading(doc, "Lost-bag Events", 2)
    _para(doc, text.lost(st))
    _picture(doc, charts.lost_chart(st, chart_dir))

    _heading(doc, "Observations & Discussion Points", 2)
    for b in text.observations(st, entries):
        _bullet(doc, b)

    _heading(doc, "Suggested Next Steps", 2)
    for b in text.next_steps(st):
        _bullet(doc, b)


def _signoff(doc) -> None:
    _para(doc)
    _para(doc, "Warm regards,")
    _para(doc, "ALSTEF Group", bold=True, colour=NAVY)
    _para(doc, "Support Services Team")


# --------------------------------------------------------------------------------------
# entry point
# --------------------------------------------------------------------------------------


def render(spec: ReportSpec, subtitle: str, entries_by_code: dict[str, list[dict]],
           bhs_by_code: dict[str, BhsStats], out_path: Path, chart_dir: Path) -> Path:
    doc = docx_base.open_report()
    _ensure_default_size(doc)
    banner = _banner_snapshot(doc)

    p = doc.add_paragraph(style="Title")
    _flat(p)
    _run(p, spec.title, size=20, bold=True, colour=NAVY, font=HEADING_FONT)
    _para(doc, subtitle, colour=SUBTITLE_GREY)
    for site in spec.sites:
        st = bhs_by_code.get(site.code)
        if st is not None and (note := text.scope_note(st)):
            _para(doc)
            _note(doc, note)

    reviews: list[tuple[SiteInfo, BhsStats]] = []
    for site in spec.sites:  # in a group, each site's "— Support Tickets" H1 is its subheading
        entries = entries_by_code.get(site.code, [])
        _tickets(doc, site, entries)
        st = bhs_by_code.get(site.code)
        if st is not None:
            _bhs_table(doc, site, st)
            reviews.append((site, st))

    for site, st in reviews:
        _review(doc, site, st, entries_by_code.get(site.code, []), banner, chart_dir)

    _signoff(doc)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
