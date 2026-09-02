"""Step 1 of `/build-reports`: work out the reporting week and the input manifest.

`discover_inputs()` scans `DropBox/PartA` and `DropBox/PartB`, classifies every file,
resolves the Monday-Sunday reporting week from the Part A filenames, runs a set of
hard-stop validations, and returns a `Manifest`. Anything wrong raises `DiscoveryError`
with a message meant to be shown straight to the user.
"""
from __future__ import annotations

import dataclasses
import datetime as dt
import re
import warnings
from pathlib import Path

import openpyxl
import yaml

from . import dates


def _load_workbook(path: Path):
    """openpyxl read-only + data-only, minus the noisy 'no default style' warning
    some third-party exports (GlideView, BagAnalytics) trigger."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        return openpyxl.load_workbook(path, read_only=True, data_only=True)

# --------------------------------------------------------------------------------------
# Filename patterns
# --------------------------------------------------------------------------------------

# Part A:  "Phonelogs - 24.08 to 30.08.2026.xlsx"  /  "Carlsource - 24.08 to 30.08.2026.xls"
_PARTA_RE = re.compile(
    r"^(?P<kind>phonelogs|carlsource)\s*-\s*"
    r"(?P<d1>\d{1,2})\.(?P<m1>\d{1,2})\s+to\s+"
    r"(?P<d2>\d{1,2})\.(?P<m2>\d{1,2})\.(?P<y2>\d{4})$",
    re.IGNORECASE,
)

# Part B GlideView:  "<Airport full name> - GlideView Reports Week Ending Sunday, <date>"
_GLIDEVIEW_RE = re.compile(
    r"^(?P<airport>.+?)\s*-\s*glideview reports\s+week ending sunday,\s*(?P<date>.+?)$",
    re.IGNORECASE,
)

# Part B BagAnalytics: week lives inside the file; the site is either an explicit
# "<CODE> - ..." filename prefix or, failing that, the sole site flagged
# `bhs_source: baganalytics` in sites.yaml.
_BAGANALYTICS_NAME_RE = re.compile(
    r"^(?:(?P<code>[A-Za-z0-9]{2,6})\s*-\s*)?.*baganalytics.*\breport\b",
    re.IGNORECASE,
)
_BAGANALYTICS_MARKER = "bagview reports | summary"
_BAGANALYTICS_DATE_RE = re.compile(
    r"last\s+7\s+days:\s*(\d{1,2}/\d{1,2}/\d{4})\s*-\s*(\d{1,2}/\d{1,2}/\d{4})",
    re.IGNORECASE,
)

_IGNORED_NAME_PREFIXES = ("~$", ".~")  # Office lock files


class DiscoveryError(Exception):
    """Raised when the DropBox inputs can't be resolved into a valid week + manifest."""


# --------------------------------------------------------------------------------------
# Result types
# --------------------------------------------------------------------------------------


@dataclasses.dataclass(frozen=True)
class PartAFile:
    path: Path
    kind: str  # "phonelogs" | "carlsource"
    start: dt.date
    end: dt.date


@dataclasses.dataclass(frozen=True)
class PartBInput:
    path: Path
    site: str
    source: str  # "glideview" | "baganalytics"
    week_ending: dt.date


@dataclasses.dataclass
class Manifest:
    start: dt.date
    end: dt.date
    phonelogs: Path
    carlsource: Path
    part_b: list[PartBInput]
    ignored: list[Path]
    root: Path

    def _rel(self, path: Path) -> str:
        try:
            return path.resolve().relative_to(self.root.resolve()).as_posix()
        except ValueError:
            return path.as_posix()

    def to_dict(self) -> dict:
        return {
            "week": {
                "start": self.start.isoformat(),
                "end": self.end.isoformat(),
                "subtitle": dates.format_subtitle(self.start, self.end),
                "filename_tag": dates.filename_tag(self.start, self.end),
            },
            "part_a": {
                "phonelogs": self._rel(self.phonelogs),
                "carlsource": self._rel(self.carlsource),
            },
            "part_b": [
                {
                    "file": self._rel(pb.path),
                    "site": pb.site,
                    "source": pb.source,
                    "week_ending": pb.week_ending.isoformat(),
                }
                for pb in self.part_b
            ],
            "ignored": [self._rel(p) for p in self.ignored],
        }


# --------------------------------------------------------------------------------------
# Part A: filename -> reporting week
# --------------------------------------------------------------------------------------


def parse_parta_week(path: Path) -> PartAFile:
    m = _PARTA_RE.match(path.stem.strip())
    if not m:
        raise DiscoveryError(
            f"{path.name!r} does not match the required Part A name "
            f"'<Phonelogs|Carlsource> - DD.MM to DD.MM.YYYY{path.suffix}'"
        )
    d1, m1 = int(m["d1"]), int(m["m1"])
    d2, m2, y2 = int(m["d2"]), int(m["m2"]), int(m["y2"])
    y1 = y2 - 1 if m1 > m2 else y2  # December -> January wrap
    try:
        start = dt.date(y1, m1, d1)
        end = dt.date(y2, m2, d2)
    except ValueError as exc:
        raise DiscoveryError(f"{path.name!r} has an invalid calendar date: {exc}") from exc
    return PartAFile(path=path, kind=m["kind"].lower(), start=start, end=end)


# --------------------------------------------------------------------------------------
# Part B: filename / contents -> (site, source, week ending)
# --------------------------------------------------------------------------------------


def parse_glideview(path: Path) -> tuple[str, dt.date]:
    m = _GLIDEVIEW_RE.match(path.stem.strip())
    if not m:
        raise DiscoveryError(f"{path.name!r} is not a recognisable GlideView filename")
    try:
        week_ending = dates.parse_english_date(m["date"])
    except ValueError as exc:
        raise DiscoveryError(
            f"{path.name!r}: can't read the 'Week Ending Sunday' date ({exc})"
        ) from exc
    return m["airport"].strip(), week_ending


def _looks_like_baganalytics_contents(path: Path) -> bool:
    try:
        wb = _load_workbook(path)
    except Exception:
        return False
    try:
        if "Summary" not in wb.sheetnames:
            return False
        ws = wb["Summary"]
        for row in ws.iter_rows(min_row=1, max_row=40, values_only=True):
            for value in row:
                if isinstance(value, str) and _BAGANALYTICS_MARKER in value.lower():
                    return True
        return False
    finally:
        wb.close()


def baganalytics_week(path: Path) -> tuple[dt.date, dt.date]:
    """Read the 'Date: LAST 7 DAYS: M/D/YYYY - M/D/YYYY' filter cell on the Summary sheet."""
    try:
        wb = _load_workbook(path)
    except Exception as exc:
        raise DiscoveryError(f"{path.name!r}: can't open BagAnalytics workbook ({exc})") from exc
    try:
        if "Summary" not in wb.sheetnames:
            raise DiscoveryError(
                f"{path.name!r}: no 'Summary' sheet - is this really a BagAnalytics export?"
            )
        ws = wb["Summary"]
        for row in ws.iter_rows(min_row=1, max_row=40, values_only=True):
            for value in row:
                if not isinstance(value, str):
                    continue
                m = _BAGANALYTICS_DATE_RE.search(value)
                if m:
                    return dates.parse_us_date(m.group(1)), dates.parse_us_date(m.group(2))
        raise DiscoveryError(
            f"{path.name!r}: couldn't find a 'Date: LAST 7 DAYS: M/D/YYYY - M/D/YYYY' "
            f"cell on the Summary sheet"
        )
    finally:
        wb.close()


# --------------------------------------------------------------------------------------
# sites.yaml
# --------------------------------------------------------------------------------------


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def load_sites(knowledge_dir: Path) -> dict[str, dict]:
    path = knowledge_dir / "sites.yaml"
    if not path.is_file():
        raise DiscoveryError(f"missing {path.as_posix()}")
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return data.get("sites", {}) or {}


def resolve_glideview_site(airport_name: str, sites: dict[str, dict]) -> str:
    token = airport_name.strip()
    if token.upper() in sites:  # allow a "<CODE> - GlideView Reports ..." prefix too
        return token.upper()
    target = _norm(token)
    for code, info in sites.items():
        candidate = info.get("airport_name")
        if candidate and _norm(candidate) == target:
            return code
    raise DiscoveryError(
        f"GlideView file airport {airport_name!r} matches no 'airport_name' in "
        f"knowledge/sites.yaml - add an entry for it"
    )


def resolve_baganalytics_site(sites: dict[str, dict]) -> str:
    matches = [code for code, info in sites.items() if info.get("bhs_source") == "baganalytics"]
    if len(matches) == 1:
        return matches[0]
    raise DiscoveryError(
        "knowledge/sites.yaml must mark exactly one site with 'bhs_source: baganalytics' "
        f"(found {len(matches)})"
    )


# --------------------------------------------------------------------------------------
# Phone-log content sanity check
# --------------------------------------------------------------------------------------

_EMAILED_ON_RE = re.compile(r"e-mailed on (\d{1,2}-[A-Za-z]{3}-\d{4})", re.IGNORECASE)


def phone_log_date_span(path: Path) -> tuple[dt.date, dt.date]:
    """Min/max call date from the 'Phone Logs' sheet ('Time Entered' + 'E-mailed on ...')."""
    try:
        wb = _load_workbook(path)
    except Exception as exc:
        raise DiscoveryError(f"{path.name!r}: can't open phone-log workbook ({exc})") from exc
    try:
        if "Phone Logs" not in wb.sheetnames:
            raise DiscoveryError(f"{path.name!r}: no 'Phone Logs' sheet")
        ws = wb["Phone Logs"]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise DiscoveryError(f"{path.name!r}: 'Phone Logs' sheet is empty") from None
        index = {
            str(name).strip().lower(): i
            for i, name in enumerate(header)
            if name is not None
        }
        time_col = index.get("time entered")
        comment_col = index.get("comments")

        found: list[dt.date] = []
        for row in rows:
            if time_col is not None and time_col < len(row):
                parsed = dates.coerce_cell_date(row[time_col])
                if parsed is not None:
                    found.append(parsed)
            if comment_col is not None and comment_col < len(row):
                cell = row[comment_col]
                if isinstance(cell, str):
                    for hit in _EMAILED_ON_RE.findall(cell):
                        try:
                            found.append(dates.parse_dmy_abbr(hit))
                        except ValueError:
                            pass
        if not found:
            raise DiscoveryError(
                f"{path.name!r}: couldn't read any call dates from 'Phone Logs' to "
                f"sanity-check the reporting week"
            )
        return min(found), max(found)
    finally:
        wb.close()


# --------------------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------------------


def _list_files(folder: Path) -> list[Path]:
    return sorted(
        p
        for p in folder.iterdir()
        if p.is_file()
        and not p.name.startswith(_IGNORED_NAME_PREFIXES)
        and p.name.lower() != ".gitkeep"
    )


def _classify_part_a(part_a_dir: Path) -> tuple[PartAFile, PartAFile]:
    matched: dict[str, PartAFile] = {}
    unmatched: list[str] = []
    for path in _list_files(part_a_dir):
        try:
            parsed = parse_parta_week(path)
        except DiscoveryError:
            unmatched.append(path.name)
            continue
        if parsed.kind in matched:
            raise DiscoveryError(
                f"more than one {parsed.kind!r} file in DropBox/PartA "
                f"({matched[parsed.kind].path.name!r} and {path.name!r})"
            )
        matched[parsed.kind] = parsed

    missing = {"phonelogs", "carlsource"} - matched.keys()
    if missing:
        extra = f" (unrecognised: {', '.join(unmatched)})" if unmatched else ""
        raise DiscoveryError(
            "DropBox/PartA is missing a "
            + " and a ".join(sorted(missing))
            + " file named '<Phonelogs|Carlsource> - DD.MM to DD.MM.YYYY.<ext>'"
            + extra
        )

    phone, carl = matched["phonelogs"], matched["carlsource"]
    if (phone.start, phone.end) != (carl.start, carl.end):
        raise DiscoveryError(
            "the two Part A files name different weeks: "
            f"{phone.path.name!r} says {phone.start}–{phone.end}, "
            f"{carl.path.name!r} says {carl.start}–{carl.end}"
        )
    return phone, carl


def _validate_week(start: dt.date, end: dt.date) -> None:
    if start.weekday() != 0:
        raise DiscoveryError(
            f"reporting week must start on a Monday - {start:%d.%m.%Y} is a "
            f"{dates.WEEKDAYS[start.weekday()]}"
        )
    if end.weekday() != 6:
        raise DiscoveryError(
            f"reporting week must end on a Sunday - {end:%d.%m.%Y} is a "
            f"{dates.WEEKDAYS[end.weekday()]}"
        )
    if (end - start).days != 6:
        raise DiscoveryError(
            f"reporting week must span exactly 7 days - {start:%d.%m.%Y} to "
            f"{end:%d.%m.%Y} is {(end - start).days + 1}"
        )


def _classify_part_b(
    part_b_dir: Path, end: dt.date, sites: dict[str, dict]
) -> tuple[list[PartBInput], list[Path]]:
    part_b: list[PartBInput] = []
    ignored: list[Path] = []

    for path in _list_files(part_b_dir):
        stem = path.stem.strip()
        if _GLIDEVIEW_RE.match(stem):
            airport, week_ending = parse_glideview(path)
            site = resolve_glideview_site(airport, sites)
            part_b.append(PartBInput(path, site, "glideview", week_ending))
            continue

        bag_match = _BAGANALYTICS_NAME_RE.match(stem)
        if bag_match or _looks_like_baganalytics_contents(path):
            code = bag_match.group("code") if bag_match and bag_match.group("code") else None
            if code:
                site = code.upper()
                if site not in sites:
                    raise DiscoveryError(
                        f"{path.name!r}: site code {site!r} in the filename is not in "
                        f"knowledge/sites.yaml"
                    )
            else:
                site = resolve_baganalytics_site(sites)
            _, wk_end = baganalytics_week(path)
            part_b.append(PartBInput(path, site, "baganalytics", wk_end))
            continue

        ignored.append(path)

    seen: dict[str, Path] = {}
    for pb in part_b:
        if pb.site in seen:
            raise DiscoveryError(
                f"two Part B files resolve to site {pb.site!r}: "
                f"{seen[pb.site].name!r} and {pb.path.name!r}"
            )
        seen[pb.site] = pb.path

    for pb in part_b:
        if pb.week_ending != end:
            raise DiscoveryError(
                f"{pb.path.name!r} covers the week ending {pb.week_ending:%d.%m.%Y}, "
                f"but the phone-log week ends {end:%d.%m.%Y} - fix the export or the "
                f"Part A files rather than blending two weeks"
            )
    return part_b, ignored


def discover_inputs(dropbox_dir: Path, knowledge_dir: Path) -> Manifest:
    dropbox_dir = Path(dropbox_dir)
    part_a_dir = dropbox_dir / "PartA"
    part_b_dir = dropbox_dir / "PartB"
    if not part_a_dir.is_dir():
        raise DiscoveryError(f"missing input folder {part_a_dir.as_posix()}")
    if not part_b_dir.is_dir():
        raise DiscoveryError(f"missing input folder {part_b_dir.as_posix()}")

    phone, carl = _classify_part_a(part_a_dir)
    start, end = phone.start, phone.end
    _validate_week(start, end)

    span_lo, span_hi = phone_log_date_span(phone.path)
    if span_lo < start or span_hi > end:
        raise DiscoveryError(
            f"{phone.path.name!r} contains calls dated {span_lo:%d.%m.%Y}–"
            f"{span_hi:%d.%m.%Y}, outside its named week {start:%d.%m.%Y}–"
            f"{end:%d.%m.%Y} - the file is probably mislabelled"
        )

    sites = load_sites(knowledge_dir)
    part_b, ignored = _classify_part_b(part_b_dir, end, sites)

    return Manifest(
        start=start,
        end=end,
        phonelogs=phone.path,
        carlsource=carl.path,
        part_b=sorted(part_b, key=lambda pb: pb.site),
        ignored=ignored,
        root=dropbox_dir.parent,
    )
